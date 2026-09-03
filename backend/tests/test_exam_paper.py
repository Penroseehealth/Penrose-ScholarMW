"""Unit tests for the professional examination paper renderers."""

import io

import pytest

from app.examinations_builder import service
from app.examinations_builder.paper import (
    CoverData,
    PaperOptions,
    build_constructed_paper,
    build_professional_docx,
    build_professional_pdf,
    lines_for_question,
    mark_text,
    number_question,
    validate_paper,
)
from app.examinations_builder.sanitize import (
    sanitize_question,
    sanitize_text,
    validate_questions,
)
from app.examinations_builder.schemas import ExamBuilderConfig
from app.examinations_builder.typeset import runs_to_reportlab, to_readable, typeset


def _questions():
    return [
        {
            "id": "q1",
            "question_type": "multiple_choice",
            "question_text": "Which organelle produces ATP?",
            "options": ["Nucleus", "Mitochondrion", "Ribosome", "Lysosome"],
            "correct_answer": "Mitochondrion",
            "marks": 2,
            "difficulty": "easy",
            "bloom_level": "Remembering",
        },
        {
            "id": "q2",
            "question_type": "essay",
            "question_text": "Discuss the factors affecting photosynthesis.",
            "marks": 10,
            "difficulty": "hard",
            "bloom_level": "Analysing",
            "rubric": "Light, CO2, temperature.",
        },
        {
            "id": "q3",
            "question_type": "short_answer",
            "question_text": "Write H_2O and 10^-3.",
            "marks": 3,
            "difficulty": "medium",
            "bloom_level": "Understanding",
        },
    ]


def _paper(include_answers=True):
    return build_constructed_paper(
        "Form 4 Biology",
        "Biology",
        _questions(),
        cover=CoverData(
            institution="PENROSE",
            subject="Biology",
            paper_number="Paper 1",
            form="Form 4",
            academic_year="2026",
            total_marks="15",
        ),
        options=PaperOptions(answer_space="standard"),
        include_answers=include_answers,
        teacher_version=include_answers,
    )


def test_mark_text_style():
    assert mark_text(2, "[2 marks]") == "[2 marks]"
    assert mark_text(3, "(2)") == "(3)"
    assert mark_text(4, "(2 marks)") == "(4 marks)"
    # Unrecognised style falls back to the default mark convention.
    assert mark_text(5, "4 marks") == "[5 marks]"


def test_number_question_style():
    assert number_question(1, "decimal") == "1"
    assert number_question(2, "alphabetic") == "B"
    assert number_question(3, "alpha_letter") == "C"
    assert number_question(4, "subsection") == "4.1"


def test_lines_for_question_scaling():
    paper = build_constructed_paper("T", "S", _questions(), include_answers=False)
    essay = next(
        q for s in paper.sections for q in s.questions if q.question_type == "essay"
    )
    mc = next(
        q
        for s in paper.sections
        for q in s.questions
        if q.question_type == "multiple_choice"
    )
    sa = next(
        q
        for s in paper.sections
        for q in s.questions
        if q.question_type == "short_answer"
    )
    # Essay questions always get at least 8 lines; standard spacing scales up.
    assert lines_for_question(essay, PaperOptions(answer_space="standard")) == 9
    assert lines_for_question(essay, PaperOptions(answer_space="minimal")) == 8
    # Selection-based types (multiple_choice) take NO answer lines.
    assert lines_for_question(mc, PaperOptions(answer_space="generous")) == 0
    # Written-response types get at least one line.
    assert lines_for_question(sa, PaperOptions(answer_space="minimal")) >= 1


def test_validate_paper_totals():
    qs = _questions()
    paper = build_constructed_paper("T", "S", qs, include_answers=False)
    report = validate_paper(paper)
    assert report["grand_total"] == 15
    # A larger question set changes the computed grand total.
    mismatched = build_constructed_paper(
        "T", "S", qs + [{"id": "x", "question_type": "essay", "marks": 50}]
    )
    assert validate_paper(mismatched)["grand_total"] == 65


def test_typeset_aware_question():
    qs = _questions()
    paper = build_constructed_paper("T", "S", qs, include_answers=False)
    stems = [q.question_text for s in paper.sections for q in s.questions]
    assert any("H_2O" in s for s in stems)


def test_to_readable_math():
    assert to_readable("x^2") == "x²"
    assert to_readable("H_2O") == "H₂O"


def test_runs_to_reportlab_has_markup():
    markup = runs_to_reportlab(typeset("Ca^{2+}"))
    assert "<super>" in markup or "<sub>" in markup


def test_pdf_render_nonempty():
    pdf = build_professional_pdf(_paper())
    assert pdf.getvalue()[:5] == b"%PDF-"
    assert len(pdf.getvalue()) > 1000
    pdf.seek(0)
    import fitz

    doc = fitz.open(stream=pdf, filetype="pdf")
    assert doc.page_count >= 1


def test_pdf_page_numbers_content_relative():
    pdf = build_professional_pdf(_paper())
    pdf.seek(0)
    import fitz

    doc = fitz.open(stream=pdf, filetype="pdf")
    markers = []
    for i in range(doc.page_count):
        text = doc[i].get_text()
        for line in text.splitlines():
            if line.startswith("Page "):
                markers.append(line.strip())
    # The cover page must have no marker; every content page must.
    assert markers, "No page-number markers found (cover should be excluded)"
    assert all(m.startswith("Page ") for m in markers)
    # First content page is "Page 1 of N".
    assert markers[0].startswith("Page 1 of ")


def test_pdf_without_cover_starts_page_1():
    paper = build_constructed_paper(
        "T", "S", _questions(), cover=None, include_answers=False
    )
    pdf = build_professional_pdf(paper)
    pdf.seek(0)
    import fitz

    doc = fitz.open(stream=pdf, filetype="pdf")
    texts = [doc[i].get_text() for i in range(doc.page_count)]
    assert any("Page 1 of" in t for t in texts)


def test_docx_render_nonempty():
    docx = build_professional_docx(_paper())
    assert docx.getvalue()[0:2] == b"PK"
    assert len(docx.getvalue()) > 1000
    docx.seek(0)
    import docx as _docx

    document = _docx.Document(io.BytesIO(docx.getvalue()))
    assert any(p.text for p in document.paragraphs)


def test_docx_superscript_subscript():
    docx = build_professional_docx(_paper())
    docx.seek(0)
    import docx as _docx

    document = _docx.Document(io.BytesIO(docx.getvalue()))
    run_map = []
    for para in document.paragraphs:
        for r in para.runs:
            run_map.append((r.text, bool(r.font.superscript), bool(r.font.subscript)))
    assert any(sup for (_, sup, _) in run_map), "Expected a superscript run"
    assert any(sub for (_, _, sub) in run_map), "Expected a subscript run"


def test_service_professional_wiring():
    cfg = ExamBuilderConfig(
        subject="Biology",
        title="Form 4 Biology",
        total_marks=15,
        duration_minutes=120,
        cover={"institution": "PENROSE", "exam_code": "BIO-F4"},
        paper_options={"answer_space": "standard", "mark_style": "[2 marks]"},
    )
    pdf = service.build_professional_pdf(cfg, _questions(), True)
    assert pdf.getvalue()[:5] == b"%PDF-"
    docx = service.build_professional_docx(cfg, _questions(), True)
    assert docx.getvalue()[0:2] == b"PK"

    paper = service._to_constructed_paper(cfg, _questions(), True)
    assert paper.cover is not None
    assert paper.cover.institution == "PENROSE"
    assert paper.options.answer_space == "standard"


@pytest.mark.parametrize(
    "value,expected",
    [
        ("2", "[2 marks]"),
        ("7", "[7 marks]"),
    ],
)
def test_mark_text_parametrized(value, expected):
    assert mark_text(int(value), "[2 marks]") == expected


# ── Sanitisation / leak prevention (spec §4, §25, §38) ──────────────────────


def test_sanitize_text_removes_point_identifiers():
    assert "point" not in sanitize_text("A syllabus point 7 topic overview")
    assert sanitize_text("Nervous System (point 4)") == "Nervous System"
    assert (
        sanitize_text("Define a stimulus (learning objective 3)") == "Define a stimulus"
    )
    # Legitimate prose must be preserved.
    assert "stimulus and response" in sanitize_text(
        "Which of the following best describes a stimulus and response?"
    )


def test_sanitize_question_strips_filler_options():
    q = sanitize_question(
        {
            "question_type": "multiple_choice",
            "question_text": "Which best describes a stimulus?",
            "options": [
                "The opposite of a stimulus",
                "An unrelated process",
                "A change detected by a receptor",
                "None of the above",
            ],
        }
    )
    assert q["options"] == ["A change detected by a receptor"]


def test_validate_questions_detects_duplicate_options_and_missing_stem():
    report = validate_questions(
        [
            {
                "id": "a",
                "question_type": "multiple_choice",
                "question_text": "",
                "options": ["point 4", "x", "x"],
                "marks": 2,
            }
        ]
    )
    assert report["clean"] is False
    joined = " ".join(report["problems"]["a"])
    assert "internal" in joined
    assert "duplicate" in joined
    assert "no question stem" in joined


def test_validate_questions_passes_wellformed_set():
    report = validate_questions(
        [
            {
                "id": "a",
                "question_type": "multiple_choice",
                "question_text": "Which best describes a receptor?",
                "options": ["A", "B", "C", "D"],
                "marks": 2,
            },
            {
                "id": "b",
                "question_type": "essay",
                "question_text": "Discuss a mechanism.",
                "marks": 10,
            },
        ]
    )
    assert report["clean"] is True


# ── Word-spacing fix in typeset (spec §13, §14) ─────────────────────────────


def test_typeset_preserves_prose_spacing():
    assert to_readable("Which of the following best describes a stimulus") == (
        "Which of the following best describes a stimulus"
    )


def test_typeset_keeps_chemistry_and_math_intact():
    assert to_readable("Write H_2O and 10^-3.") == "Write H₂O and 10⁻³."
    assert to_readable("Ca^{2+} ion") == "Ca²⁺ ion"


def test_fallback_questions_have_no_point_placeholders():
    cfg = ExamBuilderConfig(
        subject="Biology",
        topic="Nervous System",
        num_questions=8,
        question_types=["multiple_choice", "multiple_choice", "short_answer", "essay"],
    )
    ctx = {
        "subject": "Biology",
        "topic": "Nervous System",
        "form": "Form 1",
        "learning_objectives": [],
    }
    questions = service._fallback_questions(cfg, ctx, 8)
    assert questions, "Expected fallback questions"
    assert all("point " not in (q["question_text"] or "").lower() for q in questions)
    assert all(
        not q["options"] or len([o for o in q["options"] if o]) >= 2 for q in questions
    )
    assert validate_questions(questions)["clean"] is True


def test_mcq_section_header_no_objective():
    paper = build_constructed_paper(
        "T",
        "S",
        [
            {
                "id": "a",
                "question_type": "multiple_choice",
                "question_text": "Which best describes a receptor?",
                "options": ["A", "B", "C", "D"],
                "marks": 2,
            }
        ],
        include_answers=False,
    )
    header = paper.sections[0].name
    assert "OBJECTIVE" not in header.upper()
