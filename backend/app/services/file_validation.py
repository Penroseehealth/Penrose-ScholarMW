"""File upload validation and processing for ScholarAI multimodal inputs."""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass
from pathlib import Path

log = logging.getLogger(__name__)

ALLOWED_EXTENSIONS: dict[str, list[str]] = {
    "pdf": [".pdf"],
    "image": [".png", ".jpg", ".jpeg", ".webp"],
    "document": [".doc", ".docx"],
    "spreadsheet": [".xls", ".xlsx", ".csv"],
    "audio": [".mp3", ".wav", ".m4a", ".aac", ".ogg"],
}

ALL_ALLOWED = [ext for exts in ALLOWED_EXTENSIONS.values() for ext in exts]

MAX_FILE_SIZE_MB = 25
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024

MIME_MAP = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".doc": "application/msword",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xls": "application/vnd.ms-excel",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".mp3": "audio/mpeg",
    ".wav": "audio/wav",
    ".m4a": "audio/mp4",
    ".aac": "audio/aac",
    ".ogg": "audio/ogg",
}


@dataclass
class ValidatedFile:
    filename: str
    file_type: str
    mime_type: str
    size_bytes: int
    content: bytes


def classify_file(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    for file_type, extensions in ALLOWED_EXTENSIONS.items():
        if ext in extensions:
            return file_type
    return "unknown"


def validate_upload(filename: str, content: bytes) -> ValidatedFile | str:
    ext = Path(filename).suffix.lower()
    if ext not in ALL_ALLOWED:
        return (
            f"Unsupported file format: {ext}. Allowed: {', '.join(sorted(ALL_ALLOWED))}"
        )
    if len(content) > MAX_FILE_SIZE_BYTES:
        return f"File too large: {len(content) / (1024 * 1024):.1f}MB (max {MAX_FILE_SIZE_MB}MB)"
    if len(content) == 0:
        return "File is empty"
    if ext == ".pdf" and content[:4] != b"%PDF":
        return "File appears to be corrupted (invalid PDF header)"

    file_type = classify_file(filename)
    return ValidatedFile(
        filename=filename,
        file_type=file_type,
        mime_type=MIME_MAP.get(ext, "application/octet-stream"),
        size_bytes=len(content),
        content=content,
    )


async def extract_pdf_text(content: bytes) -> str:
    try:
        import PyPDF2

        reader = PyPDF2.PdfReader(io.BytesIO(content))
        pages = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                pages.append(f"[Page {i + 1}]\n{text}")
        return "\n\n".join(pages) if pages else "(No extractable text found in PDF)"
    except Exception as e:
        log.warning("PDF extraction failed: %s", e)
        return f"(PDF extraction error: {e})"


async def extract_docx_text(content: bytes) -> str:
    try:
        from docx import Document

        doc = Document(io.BytesIO(content))
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                style = para.style.name if para.style else ""
                if "Heading" in style:
                    level = style.replace("Heading ", "")
                    paragraphs.append(
                        f"{'#' * int(level) if level.isdigit() else '##'} {para.text}"
                    )
                else:
                    paragraphs.append(para.text)

        tables_text = []
        for i, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(" | ".join(cells))
            if rows:
                tables_text.append(f"\n**Table {i + 1}:**\n" + "\n".join(rows))

        result = "\n\n".join(paragraphs)
        if tables_text:
            result += "\n\n" + "\n\n".join(tables_text)
        return result if result.strip() else "(No text content found in document)"
    except Exception as e:
        log.warning("DOCX extraction failed: %s", e)
        return f"(Document extraction error: {e})"


async def extract_spreadsheet_text(content: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".csv":
            return _parse_csv(content.decode("utf-8", errors="replace"))

        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        results = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [
                str(c) if c is not None else f"Col_{i}" for i, c in enumerate(rows[0])
            ]
            data_rows = []
            for row in rows[1:51]:
                data_rows.append(
                    dict(zip(headers, [str(c) if c is not None else "" for c in row]))
                )
            results.append(
                f"**Sheet: {sheet_name}** ({len(rows) - 1} rows)\n"
                f"Columns: {', '.join(headers)}\n"
                f"Sample data (first {min(10, len(data_rows))} rows):\n"
                + "\n".join(
                    " | ".join(f"{k}: {v}" for k, v in row.items())
                    for row in data_rows[:10]
                )
            )
        wb.close()
        return "\n\n".join(results) if results else "(No data found in spreadsheet)"
    except Exception as e:
        log.warning("Spreadsheet extraction failed: %s", e)
        return f"(Spreadsheet extraction error: {e})"


def _parse_csv(text: str) -> str:
    lines = [line.strip() for line in text.strip().splitlines() if line.strip()]
    if not lines:
        return "(Empty CSV)"
    delimiter = ","
    if "\t" in lines[0]:
        delimiter = "\t"
    elif ";" in lines[0] and lines[0].count(";") > lines[0].count(","):
        delimiter = ";"
    headers = [h.strip() for h in lines[0].split(delimiter)]
    rows = []
    for line in lines[1:51]:
        vals = [v.strip() for v in line.split(delimiter)]
        rows.append(dict(zip(headers, vals)))
    result = f"**CSV Data** ({len(lines) - 1} rows)\nColumns: {', '.join(headers)}\n"
    for i, row in enumerate(rows[:10]):
        result += (
            f"Row {i + 1}: " + " | ".join(f"{k}: {v}" for k, v in row.items()) + "\n"
        )
    return result.strip()


async def extract_image_metadata(content: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    return (
        f"**Image: {filename}**\n"
        f"Format: {ext.lstrip('.').upper()}\n"
        f"Size: {len(content) / 1024:.1f} KB\n"
        f"Note: Image analysis requires vision-capable AI provider."
    )


async def extract_audio_metadata(content: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    return (
        f"**Audio: {filename}**\n"
        f"Format: {ext.lstrip('.').upper()}\n"
        f"Size: {len(content) / 1024:.1f} KB\n"
        f"Note: Audio transcription via speech-to-text service."
    )


async def extract_file_content(validated: ValidatedFile) -> str:
    if validated.file_type == "pdf":
        return await extract_pdf_text(validated.content)
    elif validated.file_type == "document":
        return await extract_docx_text(validated.content)
    elif validated.file_type == "spreadsheet":
        return await extract_spreadsheet_text(validated.content, validated.filename)
    elif validated.file_type == "image":
        return await extract_image_metadata(validated.content, validated.filename)
    elif validated.file_type == "audio":
        return await extract_audio_metadata(validated.content, validated.filename)
    return f"Unsupported file type: {validated.file_type}"
